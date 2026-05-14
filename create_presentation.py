"""Generate DiceFlow sales presentation for DISCO Corporation."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

DARK_BG = RGBColor(0x1A, 0x1A, 0x2E)
ACCENT_BLUE = RGBColor(0x00, 0x7B, 0xFF)
ACCENT_CYAN = RGBColor(0x00, 0xD4, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xCC, 0xCC, 0xCC)
DARK_GRAY = RGBColor(0x2D, 0x2D, 0x44)
GREEN = RGBColor(0x00, 0xE6, 0x76)
ORANGE = RGBColor(0xFF, 0x8C, 0x00)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

SLIDE_W = prs.slide_width
SLIDE_H = prs.slide_height


def add_dark_bg(slide):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = DARK_BG


def add_text_box(slide, left, top, width, height, text, font_size=18,
                 bold=False, color=WHITE, alignment=PP_ALIGN.LEFT,
                 font_name="Segoe UI"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    return tf


def add_bullet_slide_content(slide, bullets, left, top, width, height,
                             font_size=16, color=WHITE):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = bullet
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = "Segoe UI"
        p.space_before = Pt(8)
        p.level = 0
    return tf


def add_accent_bar(slide, top=Inches(1.8)):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.8), top, Inches(0.6), Pt(4)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT_BLUE
    shape.line.fill.background()


def slide_title(slide, title, subtitle=None):
    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(1.0),
                 title, font_size=32, bold=True, color=WHITE)
    add_accent_bar(slide, top=Inches(1.2))
    if subtitle:
        add_text_box(slide, Inches(0.8), Inches(1.4), Inches(10), Inches(0.8),
                     subtitle, font_size=16, color=LIGHT_GRAY)


# ─── SLIDE 1: COVER ─────────────────────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_dark_bg(slide)

add_text_box(slide, Inches(0.8), Inches(1.5), Inches(11), Inches(1.5),
             "DICEFLOW", font_size=60, bold=True, color=ACCENT_CYAN,
             font_name="Segoe UI Light")

add_text_box(slide, Inches(0.8), Inches(3.0), Inches(10), Inches(1.2),
             "Intelligent Dicing Sequence Optimization Software",
             font_size=28, color=WHITE)

add_text_box(slide, Inches(0.8), Inches(4.3), Inches(10), Inches(1.0),
             "Maximize throughput. Minimize blade wear. Eliminate planning errors.",
             font_size=18, color=LIGHT_GRAY)

add_accent_bar(slide, top=Inches(5.5))

add_text_box(slide, Inches(0.8), Inches(5.8), Inches(10), Inches(0.6),
             "Confidential — Prepared for DISCO Corporation",
             font_size=14, color=LIGHT_GRAY)

add_text_box(slide, Inches(0.8), Inches(6.3), Inches(10), Inches(0.6),
             "Andrea Mazzolari  |  2026",
             font_size=13, color=LIGHT_GRAY)


# ─── SLIDE 2: THE PROBLEM ────────────────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_dark_bg(slide)
slide_title(slide, "The Problem", "Manual dicing planning in high-mix production")

bullets = [
    "⚠  Process engineers spend 30–60 min per recipe planning cut sequences manually",
    "⚠  Complex multi-channel wafers (2–8 angles) multiply planning time exponentially",
    "⚠  Human errors in sequence ordering cause blade collisions, chipping, yield loss",
    "⚠  No single tool integrates CAD geometry with process parameters and sequence output",
    "⚠  Recipe changes require full re-planning — no incremental recalculation",
    "",
    "Result: Slower time-to-production, higher scrap rates, engineering bottleneck"
]
add_bullet_slide_content(slide, bullets, Inches(0.8), Inches(2.0),
                         Inches(11), Inches(5.0), font_size=17)


# ─── SLIDE 3: INTRODUCING DICEFLOW ──────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_dark_bg(slide)
slide_title(slide, "Introducing DiceFlow",
            "The first integrated dicing sequence optimizer")

bullets = [
    "✓  Desktop application for Windows (standalone, no cloud dependency)",
    "✓  Integrated CAD Sketch for visual cut definition",
    "✓  Bidirectional sync: CAD geometry ↔ data table in real time",
    "✓  Automatic multi-channel dicing sequence calculation",
    "✓  Process parameter management (blade height, feed speed, depth step)",
    "✓  Instant recalculation on any parameter change",
    "✓  Export-ready results compatible with DISCO machine recipe formats",
]
add_bullet_slide_content(slide, bullets, Inches(0.8), Inches(2.0),
                         Inches(11), Inches(5.0), font_size=17)


# ─── SLIDE 4: HOW IT WORKS ──────────────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_dark_bg(slide)
slide_title(slide, "How It Works", "Four-step workflow from design to machine recipe")

steps = [
    ("1", "DEFINE", "Draw cut lines on the integrated CAD Sketch\nor import from CSV/Excel"),
    ("2", "CONFIGURE", "Set process parameters per channel:\nblade height, feed speed, depth step, direction"),
    ("3", "CALCULATE", "One-click optimization of the full\ndicing sequence across all channels"),
    ("4", "EXPORT", "Generate machine-ready cut positions,\nsequence tables, and DFD files"),
]

for i, (num, title, desc) in enumerate(steps):
    left = Inches(0.8 + i * 3.1)
    top = Inches(2.3)

    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   left, top, Inches(2.8), Inches(4.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = DARK_GRAY
    shape.line.color.rgb = ACCENT_BLUE
    shape.line.width = Pt(1.5)

    add_text_box(slide, left + Inches(0.2), top + Inches(0.3),
                 Inches(2.4), Inches(0.6), num,
                 font_size=36, bold=True, color=ACCENT_CYAN)
    add_text_box(slide, left + Inches(0.2), top + Inches(1.0),
                 Inches(2.4), Inches(0.6), title,
                 font_size=18, bold=True, color=WHITE)
    add_text_box(slide, left + Inches(0.2), top + Inches(1.7),
                 Inches(2.4), Inches(2.2), desc,
                 font_size=13, color=LIGHT_GRAY)


# ─── SLIDE 5: CAD SKETCH INTEGRATION ────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_dark_bg(slide)
slide_title(slide, "Integrated CAD Sketch",
            "Visual cut definition with live data synchronization")

left_bullets = [
    "• Draw cut lines directly on the wafer/sample geometry",
    "• Multi-channel support with independent angles (θ)",
    "• FRONT and REAR cutting directions per channel",
    "• Real-time sync: every CAD edit updates the data table",
    "• Every table edit updates the CAD sketch",
    "• Undo/Redo with full state preservation",
    "• Channel visibility toggle for complex layouts",
]
add_bullet_slide_content(slide, left_bullets, Inches(0.8), Inches(2.0),
                         Inches(6.5), Inches(5.0), font_size=16)

shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                               Inches(7.8), Inches(2.0), Inches(4.8), Inches(4.8))
shape.fill.solid()
shape.fill.fore_color.rgb = DARK_GRAY
shape.line.color.rgb = ACCENT_BLUE
shape.line.width = Pt(1)

add_text_box(slide, Inches(8.5), Inches(3.5), Inches(3.5), Inches(2.0),
             "[ CAD SKETCH ]\n\nInteractive visual editor\nwith multi-channel\ncut line overlay",
             font_size=14, color=ACCENT_CYAN, alignment=PP_ALIGN.CENTER)


# ─── SLIDE 6: INPUT DATA & PARAMETERS ───────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_dark_bg(slide)
slide_title(slide, "Excel-Style Data Management",
            "Familiar spreadsheet interface for process parameters")

bullets = [
    "• Spreadsheet-like INPUT DATA table with Excel keyboard shortcuts",
    "• Per-cut parameters: Y position, theta, direction, blade height, feed speed, depth step",
    "• Channel-level parameter inheritance (set once, apply to all cuts)",
    "• Formula bar for precise value entry",
    "• CSV / XLSX import and export",
    "• Project save/load for full state preservation",
    "• Real-time validation with clear error messages",
]
add_bullet_slide_content(slide, bullets, Inches(0.8), Inches(2.0),
                         Inches(11), Inches(5.0), font_size=16)


# ─── SLIDE 7: CALCULATION ENGINE ────────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_dark_bg(slide)
slide_title(slide, "Optimization Engine",
            "Automatic sequence calculation with constraint satisfaction")

col1 = [
    "INPUT",
    "─────────────",
    "• Cut positions (Y, θ)",
    "• Process parameters",
    "• Sample geometry",
    "• Blade kerf / thickness",
    "• Tape thickness",
    "• Direction constraints",
]

col2 = [
    "OUTPUT",
    "─────────────",
    "• DICING SEQUENCE — ordered cuts",
    "• CUT POSITIONS — absolute coords",
    "• DICING DETAILS — full recipe table",
    "• Channel-aware sequencing",
    "• Collision-free ordering",
    "• DFD export format",
]

add_bullet_slide_content(slide, col1, Inches(0.8), Inches(2.0),
                         Inches(5.5), Inches(5.0), font_size=15)
add_bullet_slide_content(slide, col2, Inches(6.8), Inches(2.0),
                         Inches(5.5), Inches(5.0), font_size=15)

shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                               Inches(5.8), Inches(3.8), Inches(1.0), Inches(0.5))
shape.fill.solid()
shape.fill.fore_color.rgb = ACCENT_BLUE
shape.line.fill.background()


# ─── SLIDE 8: ROI / BUSINESS VALUE ──────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_dark_bg(slide)
slide_title(slide, "Business Value & ROI",
            "Quantifiable impact for production environments")

metrics = [
    ("80%", "Faster recipe\nplanning time", "From 45 min to < 10 min\nper new recipe"),
    ("0", "Sequence\nerrors", "Validated output eliminates\nblade collision risk"),
    ("15–25%", "Less blade\nwear", "Optimized cut ordering\nreduces unnecessary travel"),
    ("3×", "Faster recipe\niteration", "Instant recalculation on\nany parameter change"),
]

for i, (number, label, detail) in enumerate(metrics):
    left = Inches(0.5 + i * 3.2)
    top = Inches(2.3)

    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   left, top, Inches(2.9), Inches(4.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = DARK_GRAY
    shape.line.color.rgb = RGBColor(0x44, 0x44, 0x66)
    shape.line.width = Pt(1)

    add_text_box(slide, left + Inches(0.2), top + Inches(0.4),
                 Inches(2.5), Inches(1.0), number,
                 font_size=40, bold=True, color=GREEN)
    add_text_box(slide, left + Inches(0.2), top + Inches(1.6),
                 Inches(2.5), Inches(1.0), label,
                 font_size=16, bold=True, color=WHITE)
    add_text_box(slide, left + Inches(0.2), top + Inches(2.8),
                 Inches(2.5), Inches(1.5), detail,
                 font_size=13, color=LIGHT_GRAY)


# ─── SLIDE 9: PARTNERSHIP WITH DISCO ────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_dark_bg(slide)
slide_title(slide, "Partnership Opportunity",
            "DiceFlow as a value-add for DISCO's ecosystem")

bullets = [
    "🔹  Bundle DiceFlow with DISCO dicing saws as a premium software add-on",
    "🔹  OEM licensing: white-label or DISCO-branded version",
    "🔹  Direct export to DISCO machine recipe formats (DFD)",
    "🔹  Reduce DISCO support costs by eliminating user recipe errors",
    "🔹  Differentiator vs. competition: integrated CAD + optimization",
    "",
    "Licensing models available:",
    "   •  Per-seat perpetual license",
    "   •  Annual subscription per machine",
    "   •  Site license for high-volume fabs",
    "   •  OEM integration fee + royalty",
]
add_bullet_slide_content(slide, bullets, Inches(0.8), Inches(2.0),
                         Inches(11), Inches(5.0), font_size=16)


# ─── SLIDE 10: NEXT STEPS ───────────────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_dark_bg(slide)
slide_title(slide, "Next Steps")

bullets = [
    "1.  Live demo with your process engineering team",
    "2.  Pilot deployment at a selected DISCO customer site",
    "3.  Technical integration review (DFD export format alignment)",
    "4.  Commercial terms discussion",
    "",
    "",
    "Contact:",
    "Andrea Mazzolari",
    "andrea.mazzolari@email.com",
]
add_bullet_slide_content(slide, bullets, Inches(0.8), Inches(2.0),
                         Inches(11), Inches(5.0), font_size=18)

add_text_box(slide, Inches(0.8), Inches(6.2), Inches(10), Inches(0.6),
             "\"DiceFlow — Where precision meets productivity.\"",
             font_size=20, bold=True, color=ACCENT_CYAN,
             alignment=PP_ALIGN.LEFT)


# ─── SAVE ────────────────────────────────────────────────────────────────────
output_path = "/workspace/DiceFlow_Presentation_DISCO.pptx"
prs.save(output_path)
print(f"Presentation saved: {output_path}")
print(f"Slides: {len(prs.slides)}")
